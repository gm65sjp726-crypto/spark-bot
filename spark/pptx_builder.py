"""Сборка .pptx в формате 16:9 с оформлением по выбранной теме."""
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt


@dataclass(frozen=True)
class Theme:
    name: str
    bg: str          # фон слайдов
    title_bg: str    # фон титульного/финального слайда
    text: str        # основной текст
    title_text: str  # текст на титульном
    accent: str      # акцент (полосы, номера, маркеры)
    muted: str       # второстепенный текст
    font: str


THEMES: dict[str, Theme] = {
    "business": Theme("💼 Бизнес", "FFFFFF", "1F3A5F", "1E293B", "FFFFFF", "2F80ED", "64748B", "Calibri"),
    "dark":     Theme("🌙 Тёмная", "111827", "0B0F19", "E5E7EB", "FFFFFF", "A78BFA", "9CA3AF", "Segoe UI"),
    "minimal":  Theme("⚪ Минимализм", "FAFAFA", "FFFFFF", "111111", "111111", "111111", "737373", "Helvetica"),
    "nature":   Theme("🌿 Природа", "F4F7F2", "2D5A3D", "1F2D24", "FFFFFF", "4CAF50", "6B7F70", "Georgia"),
    "sunset":   Theme("🌅 Закат", "FFF7F0", "C2410C", "3B1D0F", "FFFFFF", "F97316", "92603F", "Trebuchet MS"),
    "ocean":    Theme("🌊 Океан", "F0F9FF", "0C4A6E", "0F2A3D", "FFFFFF", "0EA5E9", "52708A", "Verdana"),
    "tech":     Theme("🤖 Технологии", "0A0A0A", "000000", "D4D4D4", "00FFAA", "00FFAA", "8A8A8A", "Consolas"),
    "school":   Theme("🎓 Учебная", "FFFDF5", "3730A3", "1E1B4B", "FFFFFF", "F59E0B", "6366F1", "Arial"),
}

W, H = Inches(13.333), Inches(7.5)


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _fill_bg(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _rect(slide, x, y, w, h, color: str, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text: str, *, size: int, color: str, font: str,
          bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = _rgb(color)
    return box


def _fit(text: str, base: int, long_at: int, small: int) -> int:
    return small if len(text) > long_at else base


def _title_slide(prs, t: Theme, title: str, subtitle: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(s, t.title_bg)
    _rect(s, Inches(0.9), Inches(2.3), Inches(0.12), Inches(2.6), t.accent)
    _rect(s, W - Inches(3.2), Inches(-1.2), Inches(4.4), Inches(4.4), t.accent, MSO_SHAPE.OVAL)
    _text(s, Inches(1.3), Inches(2.1), Inches(9.5), Inches(2.2), title,
          size=_fit(title, 48, 50, 36), color=t.title_text, font=t.font, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _text(s, Inches(1.3), Inches(4.4), Inches(9.5), Inches(1.0), subtitle,
          size=22, color=t.title_text if t.title_bg != t.bg else t.muted, font=t.font)
    return s


def _content_slide(prs, t: Theme, idx: int, total: int, title: str, bullets: list[str], notes: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(s, t.bg)
    _rect(s, 0, 0, W, Inches(0.12), t.accent)
    _text(s, Inches(0.8), Inches(0.55), Inches(1.2), Inches(0.6), f"{idx:02d}",
          size=20, color=t.accent, font=t.font, bold=True)
    _text(s, Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.2), title,
          size=_fit(title, 36, 45, 28), color=t.text, font=t.font, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    bullets = bullets[:6]
    size = 22 if len(bullets) <= 4 and max((len(b) for b in bullets), default=0) < 90 else 18
    top, gap = Inches(2.5), min(Inches(4.4) / max(len(bullets), 1), Inches(1.0))
    for i, b in enumerate(bullets):
        y = top + gap * i
        _rect(s, Inches(0.85), y + Inches(0.14), Inches(0.16), Inches(0.16), t.accent, MSO_SHAPE.OVAL)
        _text(s, Inches(1.3), y, Inches(11.2), gap, b, size=size, color=t.text, font=t.font)

    _text(s, W - Inches(1.6), H - Inches(0.6), Inches(1.0), Inches(0.4), f"{idx} / {total}",
          size=12, color=t.muted, font=t.font, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def _final_slide(prs, t: Theme, conclusion: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(s, t.title_bg)
    _text(s, Inches(1), Inches(2.2), W - Inches(2), Inches(1.6), conclusion,
          size=_fit(conclusion, 36, 60, 28), color=t.title_text, font=t.font, bold=True,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _rect(s, W / 2 - Inches(0.8), Inches(4.05), Inches(1.6), Inches(0.08), t.accent)
    _text(s, Inches(1), Inches(4.4), W - Inches(2), Inches(0.8), "Спасибо за внимание!",
          size=22, color=t.title_text, font=t.font, align=PP_ALIGN.CENTER)
    return s


def build_pptx(deck: dict, theme_key: str, path: str) -> str:
    t = THEMES.get(theme_key, THEMES["business"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    _title_slide(prs, t, deck["title"], deck.get("subtitle", ""))
    slides = deck["slides"]
    for i, sl in enumerate(slides, 1):
        _content_slide(prs, t, i, len(slides), sl["title"], sl["bullets"], sl.get("notes", ""))
    _final_slide(prs, t, deck.get("conclusion") or deck["title"])

    prs.save(path)
    return path
